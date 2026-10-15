"""Pruebas del lector de formularios Word con documentos sintéticos (sin datos reales)."""
import io
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from utils.importar_plantilla_word import leer_plantilla_word  # noqa: E402


def _word(armar) -> bytes:
    doc = Document()
    enc = doc.sections[0].header.add_table(rows=1, cols=3, width=doc.sections[0].page_width)
    enc.rows[0].cells[1].text = "REGISTRO DE PRUEBA"
    enc.rows[0].cells[2].text = "Código: R-XX-01 Versión: 02"
    pie = doc.sections[0].footer.add_table(rows=1, cols=3, width=doc.sections[0].page_width)
    pie.rows[0].cells[0].text = "Creado Por: Área de prueba"
    pie.rows[0].cells[1].text = "Aprobado y Autorizado por: Persona de prueba"
    pie.rows[0].cells[2].text = "Fecha de Revisión: 01/01/2026"
    armar(doc)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def _bitacora(doc):
    t = doc.add_table(rows=6, cols=4)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Encargado:", "Jefe de prueba"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Frecuencia:", "Mensual"
    for celda, texto in zip(t.rows[2].cells, ("FECHA", "PRODUCTO", "LITROS", "RESPONSABLE")):
        celda.text = texto


def _checklist(doc):
    t = doc.add_table(rows=7, cols=4)
    t.rows[0].cells[0].text = "Fecha:"
    for celda, texto in zip(t.rows[1].cells, ("ITEM", "SI", "NO", "OBSERVACION")):
        celda.text = texto
    t.rows[2].cells[0].text = "CASETA"
    t.rows[3].cells[0].text = "1. ¿Hay pérdidas de agua?"
    t.rows[4].cells[0].text = "2. ¿Cables en buen estado?"
    t.rows[5].cells[0].text = "TERRENO"
    t.rows[6].cells[0].text = "3. Goteros en buen estado"


def test_encabezado_y_pie():
    p = leer_plantilla_word(_word(_bitacora), "R-XX-01 prueba.docx")
    assert p.codigo == "R-XX-01" and p.version == "02"
    assert p.nombre == "Registro de prueba"
    assert p.formato["creado_por"] == "Área de prueba"
    assert p.formato["aprobado_por"] == "Persona de prueba"
    assert p.formato["fecha_revision"] == "01/01/2026"


def test_bitacora_con_encargado_y_frecuencia():
    p = leer_plantilla_word(_word(_bitacora))
    assert p.formato["encargado"] == "Jefe de prueba" and p.formato["frecuencia"] == "Mensual"
    tabla = next(c for c in p.campos if c["tipo"] == "tabla")
    assert [(c["etiqueta"], c["tipo"]) for c in tabla["columnas"]] == [
        ("FECHA", "fecha"), ("PRODUCTO", "texto"), ("LITROS", "numero"), ("RESPONSABLE", "texto")]


def test_checklist_con_secciones():
    p = leer_plantilla_word(_word(_checklist))
    assert any(c["tipo"] == "fecha" for c in p.campos)
    checklist = next(c for c in p.campos if c["tipo"] == "checklist")
    assert checklist["con_observacion"]
    assert [(i["seccion"], i["texto"][:2]) for i in checklist["items"]] == [
        ("Caseta", "1."), ("Caseta", "2."), ("Terreno", "3.")]


def test_rellenar_sobre_el_word_original():
    from utils.rellenar_word import rellenar_word

    original = _word(_bitacora)
    p = leer_plantilla_word(original)
    tabla = next(c for c in p.campos if c["tipo"] == "tabla")
    filas = [{"fecha": "2026-10-01", "producto": f"P{k}", "litros": k, "responsable": "Ana"} for k in range(5)]
    doc = Document(io.BytesIO(rellenar_word(original, p.campos, {tabla["id"]: filas})))
    t = doc.tables[0]
    # 3 filas vacías en el Word + 2 agregadas con el mismo formato
    assert len(t.rows) == 3 + 5
    assert [c.text for c in t.rows[3].cells] == ["01-10-2026", "P0", "0", "Ana"]
    assert t.rows[-1].cells[1].text == "P4"
    assert "R-XX-01" in doc.sections[0].header.tables[0].rows[0].cells[2].text


def test_rellenar_checklist():
    from utils.rellenar_word import rellenar_word

    original = _word(_checklist)
    p = leer_plantilla_word(original)
    fecha = next(c for c in p.campos if c["tipo"] == "fecha")
    checklist = next(c for c in p.campos if c["tipo"] == "checklist")
    valores = {fecha["id"]: "2026-10-07",
               checklist["id"]: {"i1": {"r": "no", "obs": "Se repara"}, "i3": {"r": "si", "obs": ""}}}
    t = Document(io.BytesIO(rellenar_word(original, p.campos, valores))).tables[0]
    assert t.rows[0].cells[1].text == "07-10-2026"
    assert [c.text for c in t.rows[3].cells] == ["1. ¿Hay pérdidas de agua?", "", "X", "Se repara"]
    assert t.rows[6].cells[1].text == "X"


def test_codigo_planta_satelite_pk():
    """Las plantillas de planta satélite usan código PK en vez de R-OC-CE."""
    doc = Document()
    enc = doc.sections[0].header.add_table(rows=1, cols=3, width=doc.sections[0].page_width)
    enc.rows[0].cells[1].text = "REGISTRO PLANTA SATÉLITE"
    enc.rows[0].cells[2].text = "Código: PK-OC-CE-03 Versión: 01"
    _bitacora(doc)
    salida = io.BytesIO()
    doc.save(salida)
    p = leer_plantilla_word(salida.getvalue(), "PK-OC-CE-03 prueba.docx")
    assert p.codigo == "PK-OC-CE-03" and p.version == "01"
