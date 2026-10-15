"""Exporta una auditoría interna a Excel con el formato de las planillas del cliente:
una hoja por cláusula con el encabezado LISTA DE VERIFICACIÓN (protocolo, lugar, auditor, fecha, alcance y
cumplimiento), las secciones y los requisitos con SI / NO / N/A y observaciones. Al final, la hoja de
no conformidades con el plan de acción del jefe de área.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils.modelos import Auditoria, NoConformidad, Requisito

LINEA = Side(style="thin", color="999999")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
VERDE = PatternFill("solid", fgColor="E2EFDA")
GRIS = PatternFill("solid", fgColor="D9D9D9")
AJUSTE = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fecha(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _celda(hoja, fila: int, col: int, valor, negrita=False, relleno=None, alineacion=AJUSTE, borde=True):
    c = hoja.cell(row=fila, column=col, value=valor)
    c.font = Font(bold=negrita, size=10)
    c.alignment = alineacion
    if relleno:
        c.fill = relleno
    if borde:
        c.border = BORDE
    return c


def _hoja_clausula(libro, auditoria: Auditoria, clausula: str, titulo: str, requisitos: list[Requisito],
                   respuestas: dict) -> None:
    nombre = f"{clausula} {titulo}"[:31].replace("/", "-").replace(":", "")
    hoja = libro.create_sheet(nombre)
    for col, ancho in zip("ABCDEF", (12, 70, 6, 6, 6, 45)):
        hoja.column_dimensions[col].width = ancho

    conteo = {"si": 0, "no": 0, "na": 0}
    for r in requisitos:
        resp = respuestas.get(r.id)
        if resp and resp.resultado in conteo:
            conteo[resp.resultado] += 1
    evaluados = conteo["si"] + conteo["no"]

    protocolo = auditoria.version.protocolo
    _celda(hoja, 1, 2, "LISTA DE VERIFICACIÓN", negrita=True, relleno=VERDE)
    _celda(hoja, 1, 3, "CUMPLIMIENTO", negrita=True, relleno=VERDE, alineacion=CENTRO)
    hoja.merge_cells(start_row=1, start_column=3, end_row=1, end_column=5)
    _celda(hoja, 2, 2, f"PROTOCOLO {protocolo.nombre.upper()} {auditoria.version.nombre.upper()}", negrita=True)
    for col, t in zip((3, 4, 5), ("SI", "NO", "N/A")):
        _celda(hoja, 2, col, t, negrita=True, alineacion=CENTRO)
    _celda(hoja, 3, 2, f"LUGAR: {auditoria.unidad.nombre.upper()}")
    for col, k in zip((3, 4, 5), ("si", "no", "na")):
        _celda(hoja, 3, col, conteo[k], alineacion=CENTRO)
    _celda(hoja, 4, 2, f"AUDITOR: {(auditoria.auditor or '').upper()}")
    _celda(hoja, 4, 3, f"{round(conteo['si'] / evaluados * 100)} %" if evaluados else "—", alineacion=CENTRO)
    hoja.merge_cells(start_row=4, start_column=3, end_row=4, end_column=5)
    fechas = _fecha(auditoria.fecha) + (f" al {_fecha(auditoria.fecha_fin)}" if auditoria.fecha_fin else "")
    _celda(hoja, 5, 2, f"FECHA: {fechas}")
    _celda(hoja, 6, 2, f'ALCANCE: CLÁUSULA {clausula} "{titulo.upper()}"')

    fila = 8
    seccion_actual = None
    for r in requisitos:
        if r.seccion != seccion_actual:
            seccion_actual = r.seccion
            fila += 1
            _celda(hoja, fila, 1, f"{r.seccion}  {r.seccion_titulo or ''}".strip(), negrita=True, relleno=VERDE)
            hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
            if r.nivel == "fundamental":
                fila += 1
                _celda(hoja, fila, 1, "Fundamental", negrita=True)
                hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
            fila += 1
            for col, t in enumerate(("Cláusula", "Requisitos", "SI", "NO", "N/A", "OBSERVACIONES"), start=1):
                _celda(hoja, fila, col, t, negrita=True, relleno=GRIS, alineacion=CENTRO)
        fila += 1
        resp = respuestas.get(r.id)
        _celda(hoja, fila, 1, r.codigo)
        _celda(hoja, fila, 2, r.texto.split("\n\nCriterio")[0])
        for col, k in zip((3, 4, 5), ("si", "no", "na")):
            _celda(hoja, fila, col, "x" if resp and resp.resultado == k else "", alineacion=CENTRO)
        _celda(hoja, fila, 6, resp.comentario if resp else "")


def _hoja_no_conformidades(libro, auditoria: Auditoria, ncs: list[NoConformidad]) -> None:
    hoja = libro.create_sheet("No conformidades")
    columnas = ("Título auditoría", "Plazo", "Cláusula", "Pregunta", "Comentario", "Responsable",
                "Porqué 1", "Porqué 2", "Porqué 3", "Porqué 4", "Porqué 5", "Acción correctiva",
                "Acción preventiva", "Reincidencias", "Estado")
    for col, (t, ancho) in enumerate(zip(columnas, (30, 12, 10, 50, 40, 22, 25, 25, 25, 25, 25, 40, 40, 14, 12)), start=1):
        _celda(hoja, 1, col, t, negrita=True, relleno=GRIS, alineacion=CENTRO)
        hoja.column_dimensions[hoja.cell(row=1, column=col).column_letter].width = ancho
    for fila, nc in enumerate(ncs, start=2):
        porques = (nc.causa_raiz or []) + [""] * 5
        valores = (auditoria.titulo, _fecha(nc.plazo), nc.requisito.codigo, nc.requisito.texto.split("\n\nCriterio")[0],
                   nc.descripcion, nc.responsable.nombre if nc.responsable else "", *porques[:5],
                   nc.accion_correctiva or "", nc.accion_preventiva or "", nc.reincidencias or 0, nc.estado)
        for col, v in enumerate(valores, start=1):
            _celda(hoja, fila, col, v)


def exportar_auditoria(auditoria: Auditoria, clausulas: list[tuple[str, str, list[Requisito]]],
                       ncs: list[NoConformidad]) -> bytes:
    """clausulas: [(código, título, requisitos)] en el orden a exportar."""
    libro = Workbook()
    libro.remove(libro.active)
    respuestas = {r.requisito_id: r for r in auditoria.respuestas}
    for clausula, titulo, requisitos in clausulas:
        _hoja_clausula(libro, auditoria, clausula, titulo, requisitos, respuestas)
    if ncs:
        _hoja_no_conformidades(libro, auditoria, ncs)
    if not libro.worksheets:
        libro.create_sheet("Sin cláusulas")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
