"""Exporta un registro aprobado a PDF o Word recreando la base de los formularios R-OC-CE:
encabezado con logo, título, código, versión y página; encargado y frecuencia; contenido; pie con
creado por, aprobado por, fecha de revisión y dirección.
"""
import io
from datetime import date

from docx import Document as DocumentoWord
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fpdf import FPDF

from utils import almacenamiento
from utils.modelos import Registro

TIPOS_ARCHIVO = ("foto", "archivo", "firma")
GRIS = (217, 217, 217)
ETIQUETAS_OPCION = {"si": "SI", "no": "NO", "na": "N/A"}


# ---------- Datos comunes ----------

def _texto(valor, tipo: str) -> str:
    if valor is None or valor == "":
        return ""
    if tipo == "si_no":
        return "Sí" if valor else "No"
    if tipo == "fecha":
        try:
            return date.fromisoformat(str(valor)[:10]).strftime("%d-%m-%Y")
        except ValueError:
            return str(valor)
    if tipo == "numero":
        try:
            numero = float(valor)
            return str(int(numero)) if numero.is_integer() else f"{numero:.2f}".rstrip("0").replace(".", ",")
        except (TypeError, ValueError):
            return str(valor)
    return str(valor)


def _leer_archivo(clave: str | None) -> bytes | None:
    if not clave:
        return None
    try:
        return almacenamiento.leer(clave)
    except Exception:
        return None


def _imagen(valor) -> bytes | None:
    if not isinstance(valor, dict) or not valor.get("clave"):
        return None
    if not (valor.get("nombre") or "").lower().endswith((".png", ".jpg", ".jpeg")):
        return None
    return _leer_archivo(valor["clave"])


def _es_simple(campo: dict) -> bool:
    return campo["tipo"] not in ("tabla", "checklist", *TIPOS_ARCHIVO)


def _meta(registro: Registro, f: dict) -> list[tuple[str, str]]:
    """Encabezado del formulario en una sola tabla: encargado, frecuencia, datos del registro y campos simples."""
    simples = [c for c in registro.campos if _es_simple(c)]
    tiene_responsable = any(c["etiqueta"].strip().lower().rstrip(":.") == "responsable" for c in simples)
    filas = [("Encargado", f.get("encargado", "")), ("Frecuencia", f.get("frecuencia", "")),
             ("Registro", registro.titulo), ("Sucursal", registro.unidad.etiqueta)]
    if not tiene_responsable and registro.responsable:
        filas.append(("Responsable", registro.responsable.nombre))
    filas += [(c["etiqueta"], _texto(registro.valores.get(c["id"]), c["tipo"])) for c in simples]
    return [(e, v) for e, v in filas if v or e not in ("Encargado", "Frecuencia")]


def _base(registro: Registro) -> dict:
    f = registro.formato or {}
    return {
        "titulo": (registro.plantilla.nombre or registro.titulo).upper(),
        "codigo": registro.codigo or "", "version": registro.version or "01",
        "encargado": f.get("encargado", ""), "frecuencia": f.get("frecuencia", ""),
        "horizontal": f.get("orientacion") == "horizontal",
        "creado_por": f.get("creado_por", ""), "aprobado_por": f.get("aprobado_por", ""),
        "fecha_revision": f.get("fecha_revision", ""), "direccion": f.get("direccion", ""),
        "notas": f.get("notas", []), "logo": _leer_archivo(f.get("logo_clave")),
        "meta": _meta(registro, f),
        "aprobacion": (f"Aprobado en la aplicación por {registro.aprobado_por.nombre} el "
                       f"{registro.aprobado_en:%d-%m-%Y %H:%M}") if registro.aprobado_por and registro.aprobado_en else "",
    }


def _filas_checklist(campo: dict, valor) -> list[tuple[str, str, dict]]:
    """[(tipo, texto, respuesta)] con filas de sección intercaladas."""
    respuestas = valor if isinstance(valor, dict) else {}
    filas, seccion = [], None
    for it in campo.get("items", []):
        if it.get("seccion") and it["seccion"] != seccion:
            seccion = it["seccion"]
            filas.append(("seccion", seccion.upper(), {}))
        filas.append(("item", it["texto"], respuestas.get(it["id"]) or {}))
    return filas


# ---------- PDF ----------

def _latin(texto: str) -> str:
    return (texto or "").replace("–", "-").replace("·", "-").encode("latin-1", "replace").decode("latin-1")


class _PDF(FPDF):
    def __init__(self, base: dict):
        super().__init__(orientation="L" if base["horizontal"] else "P", format="Letter")
        self.base = base
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(True, margin=24)
        self.alias_nb_pages()

    def header(self):
        b, ancho = self.base, self.w - self.l_margin - self.r_margin
        x, y, alto = self.l_margin, 10, 18
        c1, c3 = ancho * 0.2, ancho * 0.24
        c2 = ancho - c1 - c3
        self.set_draw_color(0, 0, 0)
        self.rect(x, y, c1, alto)
        self.rect(x + c1, y, c2, alto)
        self.rect(x + c1 + c2, y, c3, alto)
        if b["logo"]:
            try:
                self.image(io.BytesIO(b["logo"]), x=x + 2, y=y + 1.5, h=alto - 3, keep_aspect_ratio=True, w=c1 - 4)
            except Exception:
                pass
        self.set_font("Helvetica", "B", 11)
        self.set_xy(x + c1, y)
        self.multi_cell(c2, alto / 2 if len(b["titulo"]) > 45 else alto, _latin(b["titulo"]), align="C")
        self.set_font("Helvetica", "", 8)
        for i, linea in enumerate((f"Código: {b['codigo']}", f"Versión: {b['version']}",
                                   f"Página {self.page_no()}/{{nb}}")):
            self.set_xy(x + c1 + c2 + 2, y + 1.5 + i * 5)
            self.cell(c3 - 4, 5, _latin(linea))
        self.set_y(y + alto + 4)

    def footer(self):
        b, ancho = self.base, self.w - self.l_margin - self.r_margin
        self.set_y(-20)
        self.set_font("Helvetica", "", 7)
        partes = [f"Creado por: {b['creado_por']}", f"Aprobado y autorizado por: {b['aprobado_por']}",
                  f"Fecha de revisión: {b['fecha_revision']}"]
        y = self.get_y()
        for i, t in enumerate(partes):
            self.set_xy(self.l_margin + i * ancho / 3, y)
            self.multi_cell(ancho / 3, 3.5, _latin(t), border=1, align="C")
        self.set_xy(self.l_margin, y + 9)
        self.cell(ancho, 4, _latin(b["direccion"]), align="C")


def a_pdf(registro: Registro) -> bytes:
    b = _base(registro)
    pdf = _PDF(b)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 9)

    with pdf.table(col_widths=(25, 75), first_row_as_headings=False, line_height=5) as t:
        for etiqueta, valor in b["meta"]:
            r = t.row()
            pdf.set_font("Helvetica", "B", 9)
            r.cell(_latin(etiqueta + ":"))
            pdf.set_font("Helvetica", "", 9)
            r.cell(_latin(valor))
    pdf.ln(2)
    for nota in b["notas"]:
        pdf.set_font("Helvetica", "I", 8)
        pdf.multi_cell(0, 4, _latin(nota), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)

    for c in registro.campos:
        valor = registro.valores.get(c["id"])
        if c["tipo"] == "tabla":
            columnas = c.get("columnas", [])
            filas = [[_texto(f.get(col["id"]), col["tipo"]) for col in columnas]
                     for f in (valor if isinstance(valor, list) else [])]
            pdf.set_font("Helvetica", "B", 8)
            with pdf.table(text_align="CENTER", line_height=4.5) as t:
                r = t.row()
                for col in columnas:
                    r.cell(_latin(col["etiqueta"].upper()))
                pdf.set_font("Helvetica", "", 8)
                for f in filas:
                    r = t.row()
                    for v in f:
                        r.cell(_latin(v))
            pdf.ln(2)
        elif c["tipo"] == "checklist":
            opciones = c.get("opciones", ["si", "no"])
            obs = c.get("con_observacion", True)
            anchos = (55,) + (8,) * len(opciones) + ((30,) if obs else ())
            pdf.set_font("Helvetica", "B", 8)
            with pdf.table(col_widths=anchos, line_height=4.5) as t:
                r = t.row()
                r.cell("ITEM")
                for o in opciones:
                    r.cell(ETIQUETAS_OPCION[o], align="C")
                if obs:
                    r.cell(_latin("OBSERVACIÓN"))
                for tipo, texto, resp in _filas_checklist(c, valor):
                    r = t.row()
                    if tipo == "seccion":
                        pdf.set_font("Helvetica", "B", 8)
                        r.cell(_latin(texto), colspan=len(anchos))
                        pdf.set_font("Helvetica", "", 8)
                        continue
                    pdf.set_font("Helvetica", "", 8)
                    r.cell(_latin(texto))
                    for o in opciones:
                        r.cell("X" if resp.get("r") == o else "", align="C")
                    if obs:
                        r.cell(_latin(resp.get("obs", "")))
            pdf.ln(2)
        elif c["tipo"] in TIPOS_ARCHIVO:
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(0, 5, _latin(c["etiqueta"] + ":"), new_x="LMARGIN", new_y="NEXT")
            imagen = _imagen(valor)
            if imagen:
                pdf.image(io.BytesIO(imagen), w=50 if c["tipo"] == "firma" else 80)
            elif isinstance(valor, dict):
                pdf.set_font("Helvetica", "", 9)
                pdf.cell(0, 5, _latin(f"Adjunto: {valor.get('nombre', '')}"), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)

    if b["aprobacion"]:
        pdf.set_font("Helvetica", "I", 8)
        pdf.cell(0, 5, _latin(b["aprobacion"]), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ---------- Word ----------

def _sombrear(celda, color="D9D9D9"):
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    propiedades.append(sombra)


def _texto_celda(celda, texto: str, negrita=False, tamano=8, centrado=False):
    celda.text = ""
    p = celda.paragraphs[0]
    run = p.add_run(texto)
    run.bold, run.font.size = negrita, Pt(tamano)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _campo_pagina(parrafo, instruccion: str):
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    run.font.size = Pt(8)


def a_word(registro: Registro) -> bytes:
    b = _base(registro)
    doc = DocumentoWord()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Arial", Pt(9)
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
    if b["horizontal"]:
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width, seccion.page_height = seccion.page_height, seccion.page_width
    for lado in ("left_margin", "right_margin"):
        setattr(seccion, lado, Cm(1.5))
    seccion.top_margin, seccion.bottom_margin = Cm(1.2), Cm(1.5)
    ancho = seccion.page_width - seccion.left_margin - seccion.right_margin

    # Encabezado: logo | título | código, versión y página
    enc = seccion.header.add_table(rows=1, cols=3, width=ancho)
    enc.style = "Table Grid"
    enc.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_logo, c_titulo, c_codigo = enc.rows[0].cells
    for celda, prop in ((c_logo, 0.2), (c_titulo, 0.56), (c_codigo, 0.24)):
        celda.width = int(ancho * prop)
    if b["logo"]:
        try:
            c_logo.paragraphs[0].add_run().add_picture(io.BytesIO(b["logo"]), height=Cm(1.6))
            c_logo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass
    _texto_celda(c_titulo, b["titulo"], negrita=True, tamano=12, centrado=True)
    _texto_celda(c_codigo, f"Código: {b['codigo']}")
    c_codigo.add_paragraph().add_run(f"Versión: {b['version']}").font.size = Pt(8)
    p = c_codigo.add_paragraph()
    p.add_run("Página ").font.size = Pt(8)
    _campo_pagina(p, "PAGE")
    p.add_run("/").font.size = Pt(8)
    _campo_pagina(p, "NUMPAGES")
    seccion.header.paragraphs[0].text = ""

    # Pie: creado por | aprobado por | fecha de revisión, y dirección
    pie = seccion.footer.add_table(rows=1, cols=3, width=ancho)
    pie.style = "Table Grid"
    for celda, texto in zip(pie.rows[0].cells, (f"Creado por: {b['creado_por']}",
                                                f"Aprobado y autorizado por: {b['aprobado_por']}",
                                                f"Fecha de revisión: {b['fecha_revision']}")):
        _texto_celda(celda, texto, tamano=7, centrado=True)
    direccion = seccion.footer.add_paragraph()
    direccion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = direccion.add_run(b["direccion"])
    run.font.size, run.font.color.rgb = Pt(7), RGBColor(0x55, 0x55, 0x55)

    # Encargado, frecuencia y datos del registro
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for etiqueta, valor in b["meta"]:
        celdas = t.add_row().cells
        _texto_celda(celdas[0], f"{etiqueta}:", negrita=True, tamano=9)
        _texto_celda(celdas[1], valor, tamano=9)
        celdas[0].width, celdas[1].width = int(ancho * 0.22), int(ancho * 0.78)
    for nota in b["notas"]:
        n = doc.add_paragraph()
        r = n.add_run(nota)
        r.italic, r.font.size = True, Pt(8)

    for c in registro.campos:
        valor = registro.valores.get(c["id"])
        if c["tipo"] == "tabla":
            doc.add_paragraph()
            columnas = c.get("columnas", [])
            t = doc.add_table(rows=1, cols=len(columnas))
            t.style = "Table Grid"
            for celda, col in zip(t.rows[0].cells, columnas):
                _texto_celda(celda, col["etiqueta"].upper(), negrita=True, centrado=True)
                _sombrear(celda)
            for fila in (valor if isinstance(valor, list) else []):
                for celda, col in zip(t.add_row().cells, columnas):
                    _texto_celda(celda, _texto(fila.get(col["id"]), col["tipo"]))
        elif c["tipo"] == "checklist":
            doc.add_paragraph()
            opciones = c.get("opciones", ["si", "no"])
            obs = c.get("con_observacion", True)
            columnas = ["ITEM"] + [ETIQUETAS_OPCION[o] for o in opciones] + (["OBSERVACIÓN"] if obs else [])
            t = doc.add_table(rows=1, cols=len(columnas))
            t.style = "Table Grid"
            for celda, texto in zip(t.rows[0].cells, columnas):
                _texto_celda(celda, texto, negrita=True, centrado=True)
                _sombrear(celda)
            for tipo, texto, resp in _filas_checklist(c, valor):
                celdas = t.add_row().cells
                if tipo == "seccion":
                    unida = celdas[0].merge(celdas[-1])
                    _texto_celda(unida, texto, negrita=True)
                    _sombrear(unida, "EFEFEF")
                    continue
                _texto_celda(celdas[0], texto)
                for k, o in enumerate(opciones, start=1):
                    _texto_celda(celdas[k], "X" if resp.get("r") == o else "", centrado=True)
                if obs:
                    _texto_celda(celdas[-1], resp.get("obs", ""))
            for fila in t.rows:
                fila.cells[0].width = int(ancho * (0.55 if obs else 0.8))
        elif c["tipo"] in TIPOS_ARCHIVO:
            p = doc.add_paragraph()
            p.add_run(c["etiqueta"] + ":").bold = True
            imagen = _imagen(valor)
            if imagen:
                doc.add_picture(io.BytesIO(imagen), width=Cm(5 if c["tipo"] == "firma" else 9))
            elif isinstance(valor, dict):
                doc.add_paragraph(f"Adjunto: {valor.get('nombre', '')}")

    if b["aprobacion"]:
        p = doc.add_paragraph()
        r = p.add_run(b["aprobacion"])
        r.italic, r.font.size = True, Pt(8)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
