"""Lectura de planillas: auditorías internas BRCGS y checklist oficial GLOBALG.A.P. IFA v6 GFS.

Los lectores solo interpretan y reportan; no corrigen errores en silencio.
"""
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

from utils.texto import normalizar

RE_CODIGO_ITEM = re.compile(r"^\d+(\.\d+){2,}$")
RE_SECCION = re.compile(r"^(\d+(?:\.\d+)?)\s{2,}(.+)$")
RE_FECHA = re.compile(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})")
RE_CLAUSULA_ALCANCE = re.compile(r"clausula\s+(\d+)")
RE_CONFORME = re.compile(
    r"(?<!no )\b(se evidencia cumplimiento|se observa cumplimiento|cumple|esta bien|buen estado|sin observaciones)\b")
RE_ESCAPE_MD = re.compile(r"\\([-*_.#+()\[\]])")
MAX_COLUMNAS = 12


def orden_codigo(codigo: str) -> int:
    partes = [int(p) for p in re.findall(r"\d+", codigo)][:4]
    partes += [0] * (4 - len(partes))
    return partes[0] * 1_000_000 + partes[1] * 10_000 + partes[2] * 100 + partes[3]


def _texto(valor) -> str:
    if valor is None:
        return ""
    # Los checklist oficiales traen escapes de markdown ("\- Fecha", "\*"): se dejan como texto plano
    return RE_ESCAPE_MD.sub(r"\1", str(valor)).strip()


def _comentario_conforme(comentario: str) -> bool:
    return bool(RE_CONFORME.search(normalizar(comentario)))


@dataclass
class ItemImportado:
    codigo: str
    texto: str
    seccion: str
    seccion_titulo: str
    fundamental: bool
    resultado: str | None
    comentario: str
    fila: int
    advertencias: list[str] = field(default_factory=list)


@dataclass
class HojaImportada:
    hoja: str
    lugar: str
    auditor: str
    fecha: date | None
    alcance: str
    items: list[ItemImportado] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)

    @property
    def secciones(self) -> list[str]:
        return sorted({i.codigo.split(".")[0] for i in self.items}, key=int)

    def resumen(self) -> dict:
        si = sum(1 for i in self.items if i.resultado == "si")
        no = sum(1 for i in self.items if i.resultado == "no")
        na = sum(1 for i in self.items if i.resultado == "na")
        return {
            "hoja": self.hoja, "lugar": self.lugar, "auditor": self.auditor,
            "fecha": self.fecha.isoformat() if self.fecha else None, "alcance": self.alcance,
            "secciones": self.secciones, "total": len(self.items), "si": si, "no": no, "na": na,
            "sin_respuesta": len(self.items) - si - no - na,
            "cumplimiento": round(si / (si + no) * 100, 1) if si + no else None,
            "advertencias": self.advertencias,
            "items": [
                {"codigo": i.codigo, "seccion": i.seccion, "fundamental": i.fundamental, "resultado": i.resultado,
                 "comentario": i.comentario, "texto": i.texto[:200], "fila": i.fila, "advertencias": i.advertencias}
                for i in self.items
            ],
        }


def leer_brcgs(contenido: bytes) -> list[HojaImportada]:
    libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    return [_leer_hoja_brcgs(ws) for ws in libro.worksheets]


def _leer_hoja_brcgs(ws) -> HojaImportada:
    hoja = HojaImportada(hoja=ws.title, lugar="", auditor="", fecha=None, alcance="")
    columnas: dict[str, int] = {}
    seccion, seccion_titulo, fundamental = "", "", False
    secciones_fundamentales: set[str] = set()
    vistos: set[str] = set()

    for n_fila, fila in enumerate(ws.iter_rows(max_col=MAX_COLUMNAS, values_only=True), start=1):
        celdas = [_texto(v) for v in fila]
        if not any(celdas):
            continue

        # Encabezado de la hoja
        if n_fila <= 10:
            for i, c in enumerate(celdas):
                cn = normalizar(c)
                if c == "#DIV/0!":
                    hoja.advertencias.append(
                        f"Celda {openpyxl.utils.get_column_letter(i + 1)}{n_fila} con #DIV/0! (fórmula de cumplimiento rota)")
                elif cn.startswith("lugar:"):
                    hoja.lugar = c.split(":", 1)[1].strip()
                elif cn.startswith("auditor:"):
                    hoja.auditor = c.split(":", 1)[1].strip()
                elif cn.startswith("fecha:"):
                    m = RE_FECHA.search(c)
                    if m:
                        hoja.fecha = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                elif cn.startswith("alcance:"):
                    hoja.alcance = c.split(":", 1)[1].strip()

        # Fila de encabezado de columnas: se detectan por nombre, el orden varía entre archivos
        normalizadas = [normalizar(c) for c in celdas]
        if "clausula" in normalizadas:
            nuevas = {"clausula": normalizadas.index("clausula")}
            for i, c in enumerate(normalizadas):
                if c.startswith("requisito"):
                    nuevas["requisito"] = i
                elif c == "si":
                    nuevas["si"] = i
                elif c == "no":
                    nuevas["no"] = i
                elif c in ("n/a", "na"):
                    nuevas["na"] = i
                elif c.startswith(("observacion", "comentario")):
                    nuevas["obs"] = i
            columnas.update(nuevas)
            continue

        # Marca "Fundamental" para la sección en curso
        if any(normalizar(c).startswith("fundamental") for c in celdas):
            fundamental = True
            secciones_fundamentales.add(seccion)

        # Filas de sección: "4.1     Normas externas" (pueden venir varias líneas en una celda)
        primera = celdas[0]
        if primera and not RE_CODIGO_ITEM.match(primera):
            for linea in primera.splitlines():
                m = RE_SECCION.match(linea.strip())
                if m and "edicion" not in normalizar(m.group(2)):
                    seccion, seccion_titulo = m.group(1), m.group(2).strip()
                    fundamental = seccion in secciones_fundamentales
            continue

        if "clausula" not in columnas:
            continue
        codigo = celdas[columnas["clausula"]]
        if not RE_CODIGO_ITEM.match(codigo):
            continue

        item = _leer_item(celdas, columnas, codigo, n_fila, seccion, seccion_titulo, fundamental)
        if codigo in vistos:
            item.advertencias.append("Requisito repetido en la hoja")
        vistos.add(codigo)
        hoja.items.append(item)

    # Alcance copiado de otra cláusula
    m = RE_CLAUSULA_ALCANCE.search(normalizar(hoja.alcance))
    if m and hoja.secciones and m.group(1) not in hoja.secciones:
        hoja.advertencias.append(
            f"El alcance dice cláusula {m.group(1)} pero la hoja evalúa la cláusula {', '.join(hoja.secciones)} "
            "(alcance copiado de otra planilla)")
    if not hoja.fecha:
        hoja.advertencias.append("No se encontró la fecha de la auditoría")
    if not hoja.lugar:
        hoja.advertencias.append("No se encontró el lugar")
    return hoja


def _leer_item(celdas: list[str], columnas: dict[str, int], codigo: str, n_fila: int,
               seccion: str, seccion_titulo: str, fundamental: bool) -> ItemImportado:
    advertencias: list[str] = []
    col_req = columnas.get("requisito", columnas["clausula"] + 1)
    texto = celdas[col_req] if col_req < len(celdas) else ""
    limite_texto = min(c for k, c in columnas.items() if k in ("si", "no", "na")) if "si" in columnas else len(celdas)
    if not texto:
        texto = next((c for c in celdas[columnas["clausula"] + 1:limite_texto] if c), "")

    marcas = {k: celdas[columnas[k]] for k in ("si", "no", "na") if k in columnas and celdas[columnas[k]]}
    # "x" o el propio nombre de la columna ("n/a" en N/A) cuentan como marca
    validas = [k for k, v in marcas.items() if normalizar(v) in ("x", k, "n/a" if k == "na" else k)]
    dudosas = {k: v for k, v in marcas.items() if k not in validas}
    resultado = validas[0] if len(validas) == 1 else None
    if len(validas) > 1:
        advertencias.append(f"Marcado en más de una columna ({', '.join(validas).upper()})")
    for k, v in dudosas.items():
        advertencias.append(f"Marca dudosa '{v}' en columna {k.upper()}")
    if not marcas:
        advertencias.append("Sin respuesta")

    col_obs = columnas.get("obs")
    extras = [c for i, c in enumerate(celdas) if c and col_obs is not None and i >= col_obs]
    comentario = " | ".join(extras)
    if resultado == "no" and comentario and _comentario_conforme(comentario):
        advertencias.append("Marcado 'No' pero el comentario indica cumplimiento")

    return ItemImportado(
        codigo=codigo, texto=texto or "(sin texto en la planilla)", seccion=seccion or codigo.rsplit(".", 1)[0],
        seccion_titulo=seccion_titulo, fundamental=fundamental, resultado=resultado,
        comentario=comentario, fila=n_fila, advertencias=advertencias,
    )


# ---------- GLOBALG.A.P. ----------

NIVELES_GG = {"obligacion mayor": "mayor", "obligacion menor": "menor", "recom.": "recomendacion"}


@dataclass
class PuntoControl:
    codigo: str
    seccion: str
    seccion_titulo: str
    texto: str
    nivel: str
    orden: int


def leer_globalgap(contenido: bytes) -> tuple[str, list[PuntoControl], list[str]]:
    """Devuelve (versión, puntos de control, advertencias) desde la hoja 'PI' del checklist oficial."""
    libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    advertencias: list[str] = []
    if "PI" not in libro.sheetnames:
        raise ValueError("La planilla no tiene la hoja 'PI' del checklist oficial")
    filas = list(libro["PI"].iter_rows(values_only=True))
    enc = [_texto(c) for c in filas[0]]
    idx = {nombre: enc.index(nombre) for nombre in ("Number", "P", "C", "S") if nombre in enc}
    if len(idx) < 4:
        raise ValueError("Faltan columnas esperadas en la hoja 'PI' (Number, P, C, S)")
    # El nivel se busca por contenido: los encabezados de esa planilla no coinciden con sus valores
    idx["L"] = next((i for i in range(len(enc)) if normalizar(_texto(filas[1][i])) in NIVELES_GG), -1)
    if idx["L"] < 0:
        raise ValueError("No se encontró la columna de nivel (Obligación Mayor / Menor / Recom.)")

    version = "IFA v6 GFS"
    if "Portada" in libro.sheetnames:
        for fila in libro["Portada"].iter_rows(max_row=6, max_col=1, values_only=True):
            m = re.search(r"(\d+\.\d+-GFS[_A-Z0-9]*)", _texto(fila[0]))
            if m:
                version = f"IFA v{m.group(1)}"

    puntos: list[PuntoControl] = []
    for f in filas[1:]:
        codigo = _texto(f[idx["Number"]])
        if not codigo:
            continue
        nivel_original = normalizar(_texto(f[idx["L"]]))
        nivel = NIVELES_GG.get(nivel_original)
        if not nivel:
            advertencias.append(f"{codigo}: nivel desconocido '{_texto(f[idx['L']])}'")
        s = _texto(f[idx["S"]])
        m = re.match(r"^(FV \d+(?:\.\d+)?)\s*(.*)$", s)
        seccion, titulo = (m.group(1), m.group(2).strip()) if m else (s, "")
        criterio = _texto(f[idx["C"]])
        texto = _texto(f[idx["P"]]) + (f"\n\nCriterio de cumplimiento: {criterio}" if criterio else "")
        puntos.append(PuntoControl(codigo, seccion, titulo, texto, nivel or "sin_nivel", orden_codigo(codigo)))
    puntos.sort(key=lambda p: p.orden)
    return version, puntos, advertencias
