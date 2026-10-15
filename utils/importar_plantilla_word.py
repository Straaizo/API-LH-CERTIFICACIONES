"""Lee un formulario de registro en Word (formato R-OC-CE-XX) y propone una plantilla.

Detecta: título, código y versión (encabezado), creado/aprobado/fecha de revisión y dirección (pie),
logo, encargado y frecuencia, notas, y el contenido:
- bitácora: fila(s) de títulos de columnas seguida(s) de filas vacías → campo tabla;
- checklist: ítems con columnas SI / NO (con secciones) → campo checklist;
- formulario: pares "Etiqueta: (vacío)" → campos sueltos;
- líneas de firma (JEFE COSECHA, ADMINISTRADOR…) → campos firma.
Lo que no se reconoce se informa en `advertencias` para revisarlo en el constructor.
"""
import io
import re
from dataclasses import dataclass, field

from docx import Document
from docx.table import Table

from utils.texto import normalizar

RE_CODIGO = re.compile(r"c[oó]digo\s*:?\s*([A-Z]{1,3}(?:\s*-\s*[A-Z0-9]+)+)", re.IGNORECASE)
RE_VERSION = re.compile(r"versi[oó]n\s*:?\s*(\d+)", re.IGNORECASE)
RE_SI_NO = re.compile(r"^\(?\s*si\s*/\s*no\s*\)?$")
RE_NUMERO = re.compile(r"(\bn°|\b(litros?|cantidad|ppm|ph|milimetros?|peso|punto|numero|gr|cc|kilos?|kg|volumen|horas)\b)")
RE_FIRMA = re.compile(r"\b(jefe|administrador|encargado|supervisor|firma|verificador)\b")
CLAVES_META = {"encargado": "encargado", "responsable": "encargado", "frecuencia": "frecuencia"}
TITULOS_GENERICOS = {"item", "items", "sector/ elemento", "sector/elemento", "elemento"}


@dataclass
class PlantillaDetectada:
    nombre: str = ""
    codigo: str = ""
    version: str = "01"
    formato: dict = field(default_factory=dict)
    campos: list[dict] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)
    logo: bytes | None = None
    logo_extension: str = "png"


def _limpio(t: str) -> str:
    return " ".join((t or "").split())


def _distintas(fila) -> list[str]:
    """Textos de una fila sin repetir las celdas combinadas (python-docx las devuelve repetidas)."""
    vistos, salida = set(), []
    for celda in fila.cells:
        if id(celda._tc) in vistos:
            continue
        vistos.add(id(celda._tc))
        salida.append(_limpio(celda.text))
    return salida


def _id(texto: str, usados: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", normalizar(texto)).strip("_")[:30] or "campo"
    candidato, n = base, 2
    while candidato in usados:
        candidato, n = f"{base}_{n}", n + 1
    usados.add(candidato)
    return candidato


def _tipo_columna(etiqueta: str) -> str:
    e = normalizar(etiqueta)
    if "fecha" in e:
        return "fecha"
    if RE_NUMERO.search(e):
        return "numero"
    if "?" in e or RE_SI_NO.match(e):
        return "si_no"
    return "texto"


def _titulo(texto: str) -> str:
    return texto.capitalize() if texto.isupper() else texto


def _encabezado_y_pie(doc, p: PlantillaDetectada) -> None:
    seccion = doc.sections[0]
    p.formato["orientacion"] = "horizontal" if seccion.page_width and seccion.page_height and \
        seccion.page_width > seccion.page_height else "vertical"

    textos = []
    for t in seccion.header.tables:
        for fila in t.rows:
            textos.extend(c for c in _distintas(fila) if c)
    textos.extend(_limpio(x.text) for x in seccion.header.paragraphs if x.text.strip())
    for t in textos:
        if m := RE_CODIGO.search(t):
            p.codigo = re.sub(r"\s+", "", m.group(1)).upper()
        if m := RE_VERSION.search(t):
            p.version = m.group(1).zfill(2)
    candidatos = [t for t in textos if not RE_CODIGO.search(t) and not RE_VERSION.search(t)
                  and not normalizar(t).startswith("pagina")]
    if candidatos:
        p.nombre = _titulo(max(candidatos, key=len))

    for rel in seccion.header.part.related_parts.values():
        if "image" in rel.content_type:
            p.logo = rel.blob
            p.logo_extension = rel.content_type.split("/")[-1].replace("jpeg", "jpg")
            break

    pie = []
    for t in seccion.footer.tables:
        for fila in t.rows:
            pie.extend(c for c in _distintas(fila) if c)
    pie.extend(_limpio(x.text) for x in seccion.footer.paragraphs if x.text.strip())
    for t in pie:
        n = normalizar(t)
        valor = t.split(":", 1)[1].strip() if ":" in t else ""
        if n.startswith("creado"):
            p.formato["creado_por"] = valor
        elif "aprobado" in n:
            p.formato["aprobado_por"] = valor
        elif "fecha" in n and "revision" in n:
            p.formato["fecha_revision"] = valor
        elif not valor:
            p.formato["direccion"] = t


def leer_plantilla_word(contenido: bytes, nombre_archivo: str = "") -> PlantillaDetectada:
    doc = Document(io.BytesIO(contenido))
    p = PlantillaDetectada()
    _encabezado_y_pie(doc, p)
    if not p.codigo and (m := re.match(r"([A-Z]{1,3}(?:-[A-Z0-9]+)+)", nombre_archivo.upper())):
        p.codigo = m.group(1)
    if not p.nombre:
        p.nombre = re.sub(r"^[A-Z]{1,3}(?:-[A-Z0-9]+)+\s*", "", nombre_archivo.rsplit(".", 1)[0]).strip()

    usados: set[str] = set()
    for indice, tabla in enumerate(doc.tables):
        _LectorTabla(tabla, indice, p, usados).leer()
    if not p.campos:
        p.advertencias.append("No se detectaron campos: completa la plantilla en el constructor")
    return p


def _vacia(textos: list[str]) -> bool:
    return not any(t and not t.isdigit() for t in textos)


def _con_posicion(fila) -> list[tuple[str, int]]:
    """[(texto, posición en la grilla)] de cada celda distinta de la fila."""
    vistos, salida = set(), []
    for pos, celda in enumerate(fila.cells):
        if id(celda._tc) in vistos:
            continue
        vistos.add(id(celda._tc))
        salida.append((_limpio(celda.text), pos))
    return salida


class _LectorTabla:
    """Recorre una tabla del Word. Cada campo guarda en `origen` la celda donde va su valor,
    para luego escribir los datos sobre el mismo Word."""

    def __init__(self, tabla: Table, indice: int, p: PlantillaDetectada, usados: set[str]):
        self.p, self.usados, self.indice = p, usados, indice
        self.filas = list(tabla.rows)
        self.checklist: dict | None = None
        self.cols_checklist: dict = {}
        self.seccion = ""

    def _agregar_campo(self, etiqueta: str, tipo: str = "texto", origen: dict | None = None) -> None:
        if any(normalizar(c["etiqueta"]) == normalizar(etiqueta) for c in self.p.campos):
            return
        campo = {"id": _id(etiqueta, self.usados), "etiqueta": etiqueta, "tipo": tipo, "obligatorio": True}
        if origen is not None:
            campo["origen"] = {"tabla": self.indice} | origen
        self.p.campos.append(campo)

    def _nota(self, texto: str) -> None:
        notas = self.p.formato.setdefault("notas", [])
        if texto not in notas:
            notas.append(texto)

    def leer(self) -> None:
        i = 0
        while i < len(self.filas):
            celdas = _con_posicion(self.filas[i])
            textos = [t for t, _ in celdas if t]
            normales = [normalizar(t) for t in textos]
            if not textos or _vacia(textos):
                i += 1
                continue

            es_cabecera_checklist = ("si" in normales and "no" in normales) or \
                any(RE_SI_NO.match(n) for n in normales) or "cumplimiento" in normales
            if es_cabecera_checklist:
                self._cabecera_checklist(celdas, textos, normales)
                i += 1
                continue

            if self.checklist is not None and len(textos) == 1:
                self._fila_checklist(textos[0], i, next(pos for t, pos in celdas if t))
                i += 1
                continue

            if len(textos) == 1 and len(textos[0]) > 70:
                self._nota(textos[0])
                i += 1
                continue

            if all(t.isupper() and len(t) <= 30 for t in textos) and all(RE_FIRMA.search(n) for n in normales):
                for t, pos in celdas:
                    if t:
                        self._agregar_campo(f"Firma {_titulo(t).lower()}", "firma", {"fila": i, "col": pos})
                i += 1
                continue

            if self._es_etiqueta_valor(textos, normales):
                self._etiquetas_valor(celdas, i)
                i += 1
                continue

            if len(textos) >= 3 and (fin := self._bitacora(i)) is not None:
                i = fin
                continue

            if len(textos) == 1:
                # Subtítulo de una tabla (p. ej. "USO DE JARROS")
                self._nota(_titulo(textos[0]))
                i += 1
                continue

            self.p.advertencias.append(f"Fila no reconocida: {' | '.join(textos)[:120]}")
            i += 1

    # --- checklist ---

    def _cabecera_checklist(self, celdas: list[tuple[str, int]], textos: list[str], normales: list[str]) -> None:
        if self.checklist is None:
            self.checklist = {"id": _id("checklist", self.usados), "etiqueta": "Checklist", "tipo": "checklist",
                              "obligatorio": True, "opciones": ["si", "no"], "con_observacion": False, "items": [],
                              "origen": {"tabla": self.indice, "items": {}}}
            self.p.campos.append(self.checklist)
        for t, pos in celdas:
            n = normalizar(t)
            if n == "si":
                self.cols_checklist["si"] = pos
            elif n == "no":
                self.cols_checklist["no"] = pos
            elif n in ("n/a", "na"):
                self.cols_checklist["na"] = pos
            elif RE_SI_NO.match(n) or n == "cumplimiento":
                self.cols_checklist["respuesta"] = pos
            elif "observ" in n:
                self.cols_checklist["obs"] = pos
                self.checklist["con_observacion"] = True
        primero = textos[0]
        n = normalizar(primero)
        if n not in TITULOS_GENERICOS and n not in ("si", "no", "cumplimiento") and not RE_SI_NO.match(n):
            self.seccion = _titulo(primero)

    def _fila_checklist(self, texto: str, fila: int, pos: int) -> None:
        items = self.checklist["items"]
        if texto.endswith(":"):
            self._agregar_campo(_titulo(texto.rstrip(":").strip()), "texto",
                                {"fila": fila, "col": pos, "anexar": True})
        elif texto.isupper() and "?" not in texto and len(texto) < 80:
            self.seccion = _titulo(texto)
        elif items and texto.startswith("("):
            items[-1]["texto"] += " " + texto
        else:
            item_id = f"i{len(items) + 1}"
            items.append({"id": item_id, "texto": texto, "seccion": self.seccion})
            self.checklist["origen"]["items"][item_id] = {"fila": fila, **self.cols_checklist}

    # --- etiqueta / valor ---

    @staticmethod
    def _es_etiqueta_valor(textos: list[str], normales: list[str]) -> bool:
        if any(t.endswith((":", ".")) for t in textos) or any(":" in t for t in textos):
            return True
        if normales[0].rstrip(":. ") in CLAVES_META and len(textos) <= 3:
            return True
        # "Mes", "Faena": etiqueta sola con su celda de valor vacía
        return len(textos) == 1 and len(textos[0]) <= 30 and not textos[0].isupper()

    def _etiquetas_valor(self, celdas: list[tuple[str, int]], fila: int) -> None:
        j = 0
        while j < len(celdas):
            t, pos = celdas[j]
            if not t:
                j += 1
                continue
            if ":" in t and not t.endswith(":"):
                etiqueta, valor = (x.strip() for x in t.split(":", 1))
                self._etiqueta(etiqueta, valor, {"fila": fila, "col": pos, "anexar": True})
                j += 1
                continue
            siguiente, pos_siguiente = celdas[j + 1] if j + 1 < len(celdas) else ("", None)
            es_etiqueta_siguiente = siguiente.endswith(":") or normalizar(siguiente).rstrip(":. ") in CLAVES_META
            valor = "" if es_etiqueta_siguiente else siguiente
            if pos_siguiente is not None and not es_etiqueta_siguiente:
                origen = {"fila": fila, "col": pos_siguiente}
            else:
                origen = {"fila": fila, "col": pos, "anexar": True}
            self._etiqueta(t.rstrip(":. ").strip(), valor, origen)
            j += 2 if pos_siguiente is not None and not es_etiqueta_siguiente else 1

    def _etiqueta(self, etiqueta: str, valor: str, origen: dict) -> None:
        clave = normalizar(etiqueta)
        if clave in CLAVES_META and valor:
            self.p.formato.setdefault(CLAVES_META[clave], valor)
        elif etiqueta:
            self._agregar_campo(etiqueta, "fecha" if "fecha" in clave else "texto", origen)

    # --- bitácora ---

    def _bitacora(self, inicio: int) -> int | None:
        grupo = [self.filas[inicio]]
        k = inicio + 1
        while k < len(self.filas):
            textos = [t for t in _distintas(self.filas[k]) if t]
            if _vacia(textos) or len(textos) < 2:
                break
            grupo.append(self.filas[k])
            k += 1
        if k < len(self.filas) and not _vacia([t for t in _distintas(self.filas[k]) if t]):
            return None

        # Columnas por posición en la grilla: combina títulos de varias filas ("Área" + "BF")
        ultima = grupo[-1].cells
        columnas, vistos, posiciones = [], set(), {}
        usados_col: set[str] = set()
        for pos, celda in enumerate(ultima):
            if id(celda._tc) in vistos:
                continue
            vistos.add(id(celda._tc))
            partes, es_si_no = [], False
            for fila in grupo:
                if pos >= len(fila.cells):
                    continue
                t = _limpio(fila.cells[pos].text)
                if not t:
                    continue
                if RE_SI_NO.match(normalizar(t)):
                    es_si_no = True
                elif t not in partes:
                    partes.append(t)
            if not partes:
                continue
            etiqueta = " ".join(partes[-2:])
            tipo = "si_no" if es_si_no else _tipo_columna(etiqueta)
            col_id = _id(etiqueta, usados_col)
            columnas.append({"id": col_id, "etiqueta": etiqueta, "tipo": tipo})
            posiciones[col_id] = pos

        vacias = sum(1 for f in self.filas[k:] if _vacia([t for t in _distintas(f) if t]))
        self.p.campos.append({"id": _id("registros", self.usados), "etiqueta": "Registros", "tipo": "tabla",
                              "obligatorio": True, "columnas": columnas, "filas_sugeridas": min(vacias, 31),
                              "origen": {"tabla": self.indice, "fila_inicio": k, "columnas": posiciones}})
        return len(self.filas)
