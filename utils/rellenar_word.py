"""Escribe los datos de un registro sobre el Word original del formulario (R-OC-CE-XX).

Cada campo de la plantilla trae `origen` (tabla, fila y columna de la grilla) detectado al importar el Word.
Así la descarga se ve igual que el formulario original: mismo encabezado, logo, pie, tablas y formato.
Los campos sin `origen` (agregados después en el constructor) se anexan al final del documento.
"""
import io
import re
from copy import deepcopy
from datetime import date

from docx import Document
from docx.shared import Cm, Pt

from utils import almacenamiento

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]
TIPOS_ARCHIVO = ("foto", "archivo", "firma")
RE_FRASE_FECHA = re.compile(r"CON\s+FECHA\s+DE\s+DEL\s*,", re.IGNORECASE)
RE_FRASE_CUARTEL = re.compile(r"(LIBERAR\s+EL\s+CUARTEL)", re.IGNORECASE)


def texto_valor(valor, tipo: str) -> str:
    if valor is None or valor == "":
        return ""
    if tipo == "si_no":
        return "SI" if valor else "NO"
    if tipo == "fecha":
        try:
            return date.fromisoformat(str(valor)[:10]).strftime("%d-%m-%Y")
        except ValueError:
            return str(valor)
    if tipo == "numero":
        try:
            numero = float(valor)
            return str(int(numero)) if numero.is_integer() else f"{numero:.2f}".rstrip("0").replace(".", ",")
        except (TypeError, ValueError):
            return str(valor)
    return str(valor)


def _escribir(celda, texto: str, anexar: bool = False) -> None:
    """Escribe en la celda conservando el formato del primer párrafo."""
    parrafo = celda.paragraphs[0] if not anexar else celda.paragraphs[-1]
    if anexar:
        if parrafo.runs:
            parrafo.runs[-1].text = parrafo.runs[-1].text.rstrip() + " "
        parrafo.add_run(texto)
        return
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        run = parrafo.add_run(texto)
        run.font.size = Pt(9)
    for extra in celda.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _vacia(fila) -> bool:
    return not any(c.text.strip() and not c.text.strip().isdigit() for c in fila.cells)


def _imagen(valor) -> bytes | None:
    if not isinstance(valor, dict) or not valor.get("clave"):
        return None
    if not (valor.get("nombre") or "").lower().endswith((".png", ".jpg", ".jpeg")):
        return None
    try:
        return almacenamiento.leer(valor["clave"])
    except Exception:
        return None


def _tabla(tabla, campo: dict, filas_valor: list) -> None:
    origen = campo["origen"]
    columnas = {c["id"]: c for c in campo.get("columnas", [])}
    filas = [f for f in (filas_valor or []) if any(v not in (None, "") for k, v in f.items() if not k.startswith("_"))]
    if not filas:
        return
    disponibles = []
    for i in range(origen["fila_inicio"], len(tabla.rows)):
        if not _vacia(tabla.rows[i]):
            break
        disponibles.append(i)
    if not disponibles:
        return
    # Si hay más datos que filas en el formulario, se agregan filas con el mismo formato
    ultima = tabla.rows[disponibles[-1]]._tr
    while len(disponibles) < len(filas):
        nueva = deepcopy(ultima)
        for t in nueva.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
            t.text = ""
        ultima.addnext(nueva)
        ultima = nueva
        disponibles.append(disponibles[-1] + 1)
    for fila_valor, indice in zip(filas, disponibles):
        fila = tabla.rows[indice]
        for col_id, pos in origen["columnas"].items():
            if col_id in columnas and pos < len(fila.cells):
                _escribir(fila.cells[pos], texto_valor(fila_valor.get(col_id), columnas[col_id]["tipo"]))


def _checklist(tabla, campo: dict, respuestas: dict) -> None:
    etiquetas = {"si": "SI", "no": "NO", "na": "N/A"}
    for item_id, o in campo["origen"].get("items", {}).items():
        resp = (respuestas or {}).get(item_id) or {}
        if o["fila"] >= len(tabla.rows):
            continue
        celdas = tabla.rows[o["fila"]].cells
        r = resp.get("r")
        if r and r in o and o[r] < len(celdas):
            _escribir(celdas[o[r]], "X")
        elif r and "respuesta" in o and o["respuesta"] < len(celdas):
            _escribir(celdas[o["respuesta"]], etiquetas.get(r, r))
        if resp.get("obs") and "obs" in o and o["obs"] < len(celdas):
            _escribir(celdas[o["obs"]], resp["obs"])


def _frase(tabla, fila: int, datos: dict) -> None:
    """R-OC-CE-10: completa "CON FECHA __ DE __ DEL __, ... LIBERAR EL CUARTEL __" con fecha, cuartel y especie."""
    celda = tabla.rows[fila].cells[0]
    for parrafo in celda.paragraphs:
        texto = parrafo.text
        if not RE_FRASE_FECHA.search(texto):
            continue
        try:
            f = date.fromisoformat(str(datos.get("fecha"))[:10])
            texto = RE_FRASE_FECHA.sub(f"CON FECHA {f.day} DE {MESES[f.month - 1].upper()} DEL {f.year},", texto)
        except ValueError:
            pass
        cuartel = " ".join(x for x in (str(datos.get("cuartel") or ""),
                                        f"({datos['especie']})" if datos.get("especie") else "") if x)
        if cuartel:
            texto = RE_FRASE_CUARTEL.sub(lambda m: f"{m.group(1)} {cuartel.upper()}", texto, count=1)
        if parrafo.runs:
            parrafo.runs[0].text = texto
            for run in parrafo.runs[1:]:
                run.text = ""
        return


def _anexar_adicionales(doc, campos: list[dict], valores: dict) -> None:
    if not campos:
        return
    doc.add_paragraph()
    titulo = doc.add_paragraph()
    titulo.add_run("Información registrada en la aplicación").bold = True
    for c in campos:
        valor = valores.get(c["id"])
        if c["tipo"] == "tabla":
            columnas = c.get("columnas", [])
            doc.add_paragraph().add_run(c["etiqueta"]).bold = True
            t = doc.add_table(rows=1, cols=len(columnas))
            t.style = "Table Grid"
            for celda, col in zip(t.rows[0].cells, columnas):
                celda.text = col["etiqueta"]
            for fila in (valor if isinstance(valor, list) else []):
                for celda, col in zip(t.add_row().cells, columnas):
                    celda.text = texto_valor(fila.get(col["id"]), col["tipo"])
        elif c["tipo"] == "checklist":
            doc.add_paragraph().add_run(c["etiqueta"]).bold = True
            for it in c.get("items", []):
                resp = (valor or {}).get(it["id"]) or {}
                doc.add_paragraph(f"{it['texto']}: {(resp.get('r') or '').upper()} {resp.get('obs', '')}".strip())
        elif c["tipo"] in TIPOS_ARCHIVO:
            imagen = _imagen(valor)
            p = doc.add_paragraph()
            p.add_run(f"{c['etiqueta']}: ").bold = True
            if imagen:
                p.add_run().add_picture(io.BytesIO(imagen), height=Cm(1.5))
            elif isinstance(valor, dict):
                p.add_run(valor.get("nombre", ""))
        else:
            p = doc.add_paragraph()
            p.add_run(f"{c['etiqueta']}: ").bold = True
            p.add_run(texto_valor(valor, c["tipo"]))


def rellenar_word(original: bytes, campos: list[dict], valores: dict) -> bytes:
    doc = Document(io.BytesIO(original))
    tablas = doc.tables
    sin_origen = []
    frases: dict[tuple[int, int], dict] = {}
    for campo in campos:
        origen = campo.get("origen")
        valor = valores.get(campo["id"])
        if not origen or origen.get("tabla", 0) >= len(tablas):
            sin_origen.append(campo)
            continue
        tabla = tablas[origen["tabla"]]
        try:
            if origen.get("frase"):
                frases.setdefault((origen["tabla"], origen["fila"]), {})[origen["frase"]] = valor
            elif campo["tipo"] == "tabla":
                _tabla(tabla, campo, valor)
            elif campo["tipo"] == "checklist":
                _checklist(tabla, campo, valor)
            elif origen["fila"] < len(tabla.rows):
                celda = tabla.rows[origen["fila"]].cells[origen["col"]]
                if campo["tipo"] in TIPOS_ARCHIVO:
                    imagen = _imagen(valor)
                    if imagen:
                        celda.add_paragraph().add_run().add_picture(io.BytesIO(imagen), height=Cm(1.2))
                    elif isinstance(valor, dict) and valor.get("nombre"):
                        _escribir(celda, valor["nombre"], anexar=True)
                elif valor not in (None, ""):
                    _escribir(celda, texto_valor(valor, campo["tipo"]), anexar=bool(origen.get("anexar")))
        except (IndexError, KeyError, ValueError):
            sin_origen.append(campo)
    for (indice, fila), datos in frases.items():
        _frase(tablas[indice], fila, datos)
    _anexar_adicionales(doc, [c for c in sin_origen if valores.get(c["id"]) not in (None, "", [], {})], valores)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
